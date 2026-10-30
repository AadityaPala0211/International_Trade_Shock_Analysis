"""Companion-workbook metadata and conservative OECD ICIO provenance audit."""
from pathlib import Path
import re
from openpyxl import load_workbook

OECD_ICIO_DATASET_URL = "https://www.oecd.org/en/data/datasets/inter-country-input-output-tables.html"
KEYWORDS = ("oecd", "icio", "inter-country", "input-output", "edition",
            "release", "2025", "2019", "2022", "sml", "version")


def inspect_icio_readme(path, preview_rows=30, preview_columns=12, max_evidence_cells=250):
    """Return workbook properties, a bounded preview and whole-workbook provenance hits.

    The full workbook is scanned for provenance-related cell values. Only a
    capped number of matching cells are returned to keep metadata compact.
    """
    path = Path(path)
    if not path.exists():
        return {"available": False, "path": str(path), "reason": "file_not_found"}
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        props = wb.properties
        sheets, evidence, workbook_nonempty_cells = [], [], 0
        for ws in wb.worksheets:
            preview = []
            for row in ws.iter_rows(min_row=1, max_row=min(preview_rows, ws.max_row),
                                    min_col=1, max_col=min(preview_columns, ws.max_column),
                                    values_only=True):
                values = [None if value is None else str(value) for value in row]
                if any(value is not None for value in values):
                    preview.append(values)
            sheets.append({"name": ws.title, "max_row": ws.max_row,
                           "max_column": ws.max_column, "preview": preview})

            for row_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                for column_index, value in enumerate(row, start=1):
                    if value is None:
                        continue
                    workbook_nonempty_cells += 1
                    text_value = str(value).strip()
                    lower_value = text_value.lower()
                    if any(keyword in lower_value for keyword in KEYWORDS):
                        if len(evidence) < int(max_evidence_cells):
                            evidence.append({
                                "sheet": ws.title, "cell": f"{_column_letters(column_index)}{row_index}",
                                "value": text_value[:1000],
                            })
        result = {
            "available": True, "path": str(path), "sheet_names": wb.sheetnames,
            "properties": {
                "title": props.title, "subject": props.subject, "creator": props.creator,
                "description": props.description,
                "created": str(props.created) if props.created else None,
                "modified": str(props.modified) if props.modified else None,
            },
            "sheets": sheets,
            "whole_workbook_scan": {
                "completed": True, "nonempty_cells_scanned": workbook_nonempty_cells,
                "keyword_evidence_cells_returned": len(evidence),
                "keyword_evidence_truncated": len(evidence) >= int(max_evidence_cells),
                "keyword_evidence": evidence,
            },
        }
        wb.close()
        return result
    except Exception as exc:
        return {"available": True, "path": str(path),
                "inspection_error": f"{type(exc).__name__}: {exc}"}


def _column_letters(column):
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def summarize_icio_provenance(workbook_info, loaded_csvs):
    """Separate official release context from what the local files actually prove."""
    text_parts = []
    if workbook_info.get("available") and not workbook_info.get("inspection_error"):
        scan = workbook_info.get("whole_workbook_scan", {})
        for item in scan.get("keyword_evidence", []):
            text_parts.append(item.get("value", ""))
        for sheet in workbook_info.get("sheets", []):
            text_parts.append(sheet.get("name", ""))
            for row in sheet.get("preview", []):
                text_parts.extend("" if value is None else str(value) for value in row)
        text_parts.extend(str(v) for v in workbook_info.get("properties", {}).values() if v is not None)
    companion_text = " ".join(text_parts)
    lower = companion_text.lower()
    mentions_icio = "icio" in lower or "inter-country input-output" in lower
    mentions_2025 = bool(re.search(r"(?<!\d)2025(?!\d)", companion_text))
    mentions_oecd = "oecd" in lower
    if mentions_icio and mentions_2025 and mentions_oecd:
        companion_status = "companion_workbook_mentions_oecd_icio_2025"
    elif mentions_icio and mentions_2025:
        companion_status = "companion_workbook_mentions_icio_2025"
    elif mentions_2025:
        companion_status = "companion_workbook_mentions_2025_but_not_explicitly_icio"
    else:
        companion_status = "companion_workbook_does_not_establish_2025_icio_edition"

    files = []
    for item in loaded_csvs:
        files.append({
            "path": item.get("source_path"), "sha256": item.get("source_sha256"),
            "shape": list(item.get("Z").shape) if item.get("Z") is not None else None,
            "node_count": item.get("n_nodes"),
        })
    return {
        "local_companion_workbook_status": companion_status,
        "companion_workbook_mentions_oecd": mentions_oecd,
        "companion_workbook_mentions_icio": bool(mentions_icio),
        "companion_workbook_mentions_2025": bool(mentions_2025),
        "workbook_whole_scan_completed": bool(workbook_info.get("whole_workbook_scan", {}).get("completed")),
        "csv_file_checksums_and_dimensions": files,
        "what_is_verified_locally": (
            "The selected CSV files were parsed, their dimensions and SHA-256 checksums recorded, "
            "and their detailed node universes were aligned."
        ),
        "what_is_not_independently_proven": (
            "Even a full companion-workbook scan cannot by itself prove which OECD download archive "
            "each individual CSV came from. Retain the original OECD download URL/archive or source "
            "manifest for file-level provenance."
        ),
        "official_oecd_reference": OECD_ICIO_DATASET_URL,
        "official_reference_context": (
            "The OECD ICIO dataset page lists a 2025 regular ICIO release covering 1995-2022. "
            "This supports using 2019 and 2022 from that release in principle, but does not independently "
            "identify the provenance of these local CSV files."
        ),
    }
